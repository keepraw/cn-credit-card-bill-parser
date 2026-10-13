"""Legacy migration identity audit (Task A-4.1). Read-only dry-run.

Reads legacy exports and golden fixtures, matches on BUSINESS keys only
(native IDs are never compared: algorithms differ by design), and reports
exact / suspected / orphan counts plus fingerprint-impact estimates.

Never writes databases, never migrates data, never changes identity logic.
"""

from __future__ import annotations

import argparse
import glob
import json
import re
import sys
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from ccparser.golden import (  # noqa: E402
    BANK_TO_ID,
    GOLDEN_SCHEMA,
    currency_exponent,
    normalize_description,
)

FLOAT_DUST_TOLERANCE = 1e-6

#: Chinese bank names (incl. observed variants) -> BankId. Mirrors the Web
#: legacy adapter BANK_MAP; unknown names are reported, never guessed.
BANK_MAP = {
    "中国银行": "boc",
    "交通银行": "bocom",
    "工商银行": "icbc",
    "中国工商银行": "icbc",
    "农业银行": "abc",
    "中国农业银行": "abc",
    "建设银行": "ccb",
    "中国建设银行": "ccb",
    "中信银行": "citic",
}
BANK_MAP.update({v: v for v in ("boc", "bocom", "icbc", "ccb", "abc", "citic", "unknown")})

CURRENCY_MAP = {"RMB": "CNY", "CHN": "CNY", "人民币": "CNY", "": "CNY"}

#: SIMULATION port of web classify.ts TYPE_TOKENS (order-sensitive) used ONLY
#: to estimate legacy-vs-native type forks. Exact proof lives in the TS
#: dry-run test, which runs the real functions.
TYPE_TOKENS: tuple[tuple[str, str], ...] = (
    (r"^(?:分期利息|分期手续费)", "interest"),
    (r"^(?:分期扣款|每期入账|分期入账)", "installment_payment"),
    (r"^(?:分期(?:消费)?原始|商户分期(?!扣款)|消费\s*商户分期)", "installment_original"),
    (r"^(?:信用卡还款|还款|转账还款|自助还款|网上还款)", "payment"),
    (r"^(?:退货|退款|退汇)", "refund"),
    (r"^(?:刷卡金|返现|返还|积分兑换|优惠)", "cashback"),
    (r"^(?:手续费|年费|服务费|短信服务费|短信费|违约金)", "fee"),
    (r"^(?:利息|透支利息|计息)", "interest"),
    (r"^(?:取现|预借现金)", "cash_advance"),
    (r"^(?:消费|购物|支付)", "purchase"),
    (r"^(?:调整|冲正|其他)", "adjustment"),
)


def infer_legacy_type(description: str, direction: str) -> str:
    """Mirror of inferTransactionType('', description, direction)."""
    head = (description or "").strip()
    for pattern, tx_type in TYPE_TOKENS:
        if re.match(pattern, head):
            return tx_type
    return "purchase" if direction == "debit" else "payment"


def infer_native_type(bank_token: str, description: str, direction: str) -> str:
    """Mirror of inferTransactionType(bankToken, description, direction)."""
    head = (bank_token or "").strip() or (description or "").strip()
    for pattern, tx_type in TYPE_TOKENS:
        if re.match(pattern, head):
            return tx_type
    return "purchase" if direction == "debit" else "payment"


def snap_float_to_minor(value: float, currency: str) -> int | None:
    exp = currency_exponent(currency)
    scaled = value * (10**exp)
    snapped = round(scaled)
    if abs(scaled - snapped) > FLOAT_DUST_TOLERANCE:
        return None
    return int(snapped)


def str_to_minor(text: str, currency: str) -> int | None:
    cleaned = text.strip().replace(",", "").lstrip("¥￥")
    digits = re.sub(r"[^\d.\-]", "", cleaned)
    try:
        amount = Decimal(digits) if digits not in ("", "-", ".") else None
    except InvalidOperation:
        return None
    if amount is None:
        return None
    exp = currency_exponent(currency)
    return int((amount * (10**exp)).to_integral_value())


def iso_of(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20000 <= value <= 80000:
        base = date(1899, 12, 30)
        try:
            return base.fromordinal(base.toordinal() + int(value)).isoformat()
        except ValueError:
            return None
    text = str(value).strip()
    if not text:
        return None
    match = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?", text)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
        except ValueError:
            return None
    return None


def business_key(row: dict) -> str:
    return "|".join(str(row[k]) for k in (
        "bank", "cardLast4", "transactionDate", "postingDate",
        "settlementAmountMinor", "currency", "normalizedDescription",
    ))


def loose_key(row: dict) -> str:
    return "|".join(str(row[k]) for k in (
        "bank", "cardLast4", "transactionDate", "settlementAmountMinor", "currency",
    ))


def load_standard_sheet(path: Path) -> tuple[list[str], list[dict], list[dict]]:
    """Parse a 标准字段 sheet into audit rows. Returns (header, rows, rejects)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    candidates = [name for name in wb.sheetnames if "标准字段" in name]
    if not candidates:
        return [], [], [{"row": 0, "reason": "no_standard_sheet", "sheets": wb.sheetnames}]
    ws = wb[candidates[0]]
    header = [(c.value if c.value is not None else "") for c in ws[1]]
    rows: list[dict] = []
    rejects: list[dict] = []
    for lineno, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        cell = dict(zip(header, values))
        bank_raw = str(cell.get("bank") or "").strip()
        bank = BANK_MAP.get(bank_raw)
        if bank is None:
            rejects.append({"row": lineno, "reason": "unknown_bank", "value": bank_raw})
            continue
        digits = re.sub(r"\D", "", str(cell.get("card_last4") or ""))
        card = digits[-4:].rjust(4, "0") if digits else ""
        if not card:
            rejects.append({"row": lineno, "reason": "missing_card"})
            continue
        tran = iso_of(cell.get("transaction_date"))
        post = iso_of(cell.get("posting_date"))
        if not tran and not post:
            rejects.append({"row": lineno, "reason": "unparseable_dates"})
            continue
        tran = tran or post
        post = post or tran
        currency = CURRENCY_MAP.get(str(cell.get("settlement_currency") or "").strip().upper(), None)
        if currency is None:
            currency = str(cell.get("settlement_currency") or "CNY").strip().upper() or "CNY"
        raw_amount = cell.get("settlement_amount")
        if isinstance(raw_amount, (int, float)) and not isinstance(raw_amount, bool):
            minor = snap_float_to_minor(float(raw_amount), currency)
        else:
            minor = str_to_minor(str(raw_amount) if raw_amount is not None else "", currency)
        if minor is None:
            rejects.append({"row": lineno, "reason": "unparseable_amount", "value": str(raw_amount)[:40]})
            continue
        description = re.sub(r"　", " ", str(cell.get("description") or "")).strip()
        description = re.sub(r"\s+", " ", description)
        if not description:
            rejects.append({"row": lineno, "reason": "missing_description"})
            continue
        direction = "credit" if minor < 0 else "debit"
        legacy_type = infer_legacy_type(description, direction)
        rows.append({
            "bank": bank, "bankNative": bank_raw, "cardLast4": card,
            "transactionDate": tran, "postingDate": post,
            "settlementAmountMinor": minor, "currency": currency,
            "direction": direction, "description": description,
            "normalizedDescription": normalize_description(description),
            "legacyType": legacy_type,
            "sourceRow": lineno,
        })
    return header, rows, rejects


def load_golden_rows(golden_root: Path) -> list[dict]:
    rows: list[dict] = []
    for python_path in sorted(glob.glob(str(golden_root / "*" / "*" / "python.json"))):
        doc = json.loads(Path(python_path).read_text(encoding="utf-8"))
        for tx in doc.get("transactions", []):
            rows.append({
                "bank": tx.get("bank"), "cardLast4": tx.get("cardLast4"),
                "transactionDate": tx.get("transactionDate"), "postingDate": tx.get("postingDate"),
                "settlementAmountMinor": tx.get("amountMinor"), "currency": tx.get("currency"),
                "direction": tx.get("direction"), "description": tx.get("description"),
                "normalizedDescription": tx.get("normalizedDescription"),
                "goldenFile": python_path,
            })
    return rows


def load_fresh_parse_rows(golden_root: Path, processed_dir: Path) -> tuple[list[dict], list[str]]:
    """Re-parse golden inputs in-memory with desensitization OFF.

    Measures true Python->Python business-key stability (committed goldens are
    desensitized, so they can never exact-match raw legacy text). Nothing is
    written; files are only read.
    """
    from ccparser.dedupe import assign_ids
    from ccparser.extractors import extract_source
    from ccparser.golden import normalize_python_statement
    from ccparser.hash_utils import file_sha256
    from ccparser.parsers.registry import parse_statement

    hash_to_path: dict[str, Path] = {}
    for candidate in processed_dir.iterdir():
        if candidate.is_file() and candidate.name != ".gitkeep":
            try:
                hash_to_path[file_sha256(candidate)] = candidate
            except OSError:
                continue
    rows: list[dict] = []
    missing: list[str] = []
    for meta_path in sorted(glob.glob(str(golden_root / "*" / "*" / "input.meta.json"))):
        meta = json.loads(Path(meta_path).read_text(encoding="utf-8"))
        source = hash_to_path.get(meta.get("sha256", ""))
        if source is None:
            missing.append(meta_path)
            continue
        statement = parse_statement(extract_source(source))
        assign_ids(statement)
        doc = normalize_python_statement(statement, desensitize=False)
        for tx in doc["transactions"]:
            rows.append({
                "bank": tx["bank"], "cardLast4": tx["cardLast4"],
                "transactionDate": tx["transactionDate"], "postingDate": tx["postingDate"],
                "settlementAmountMinor": tx["amountMinor"], "currency": tx["currency"],
                "direction": tx["direction"], "description": tx["description"],
                "normalizedDescription": tx["normalizedDescription"],
                "goldenFile": meta_path,
            })
    return rows, missing


def match_rows(legacy_rows: list[dict], golden_rows: list[dict]) -> dict:
    golden_by_key: dict[str, list[int]] = {}
    for index, row in enumerate(golden_rows):
        golden_by_key.setdefault(business_key(row), []).append(index)
    golden_loose: dict[str, list[int]] = {}
    for index, row in enumerate(golden_rows):
        golden_loose.setdefault(loose_key(row), []).append(index)

    exact = suspected = orphan = 0
    suspected_examples: list[dict] = []
    for row in legacy_rows:
        key = business_key(row)
        if key in golden_by_key:
            exact += 1
            continue
        loose_peers = golden_loose.get(loose_key(row), [])
        if loose_peers:
            suspected += 1
            if len(suspected_examples) < 5:
                suspected_examples.append({
                    "legacyRow": row["sourceRow"],
                    "legacyDescription": row["description"][:60],
                    "goldenPeers": [golden_rows[i]["description"][:60] for i in loose_peers[:3]],
                })
        else:
            orphan += 1
    return {
        "exact": exact, "suspected": suspected, "orphan": orphan,
        "total": len(legacy_rows), "suspectedExamples": suspected_examples,
    }


def self_collisions(rows: list[dict]) -> dict:
    by_key: Counter = Counter(business_key(r) for r in rows)
    exact_dupes = sum(count - 1 for count in by_key.values() if count > 1)
    by_loose: dict[str, set[str]] = {}
    for row in rows:
        by_loose.setdefault(loose_key(row), set()).add(row["normalizedDescription"])
    suspected_groups = sum(1 for descs in by_loose.values() if len(descs) > 1)
    return {
        "rows": len(rows),
        "exactDuplicateRows": exact_dupes,
        "suspectedLooseGroups": suspected_groups,
    }


def cmd_audit(args: argparse.Namespace) -> int:
    standard_path = Path(args.standard)
    header, legacy_rows, rejects = load_standard_sheet(standard_path)
    golden_rows = load_golden_rows(Path(args.golden_dir))
    matching = match_rows(legacy_rows, golden_rows)
    fresh_rows, fresh_missing = load_fresh_parse_rows(Path(args.golden_dir), Path("processed"))
    fresh_matching = match_rows(legacy_rows, fresh_rows)
    collisions = self_collisions(legacy_rows)

    type_forks = sum(
        1 for r in legacy_rows
        if r["legacyType"] in ("installment_payment", "installment_original", "fee", "interest", "cash_advance")
    )
    generatable = len(legacy_rows)
    reject_reasons = Counter(r["reason"] for r in rejects)

    bank_counts = Counter(r["bank"] for r in legacy_rows)
    dates = [r["transactionDate"] for r in legacy_rows if r["transactionDate"]]
    report = {
        "schema": GOLDEN_SCHEMA,
        "inputs": {
            "standard": {
                "file": str(standard_path), "header": header,
                "dataRows": generatable + len(rejects),
                "parsedRows": generatable, "rejects": len(rejects),
                "rejectReasons": dict(reject_reasons),
            },
            "goldenRows": len(golden_rows),
        },
        "legacyProfile": {
            "banks": dict(bank_counts),
            "dateMin": min(dates) if dates else None,
            "dateMax": max(dates) if dates else None,
        },
        "businessKeyMatching": matching,
        "freshParseMatching": {
            **{k: v for k, v in fresh_matching.items() if k != "suspectedExamples"},
            "freshRows": len(fresh_rows),
            "inputsNotFound": fresh_missing,
            "note": ("Legacy export rows vs fresh undesensitized re-parses of the same "
                     "bills: isolates pipeline stability from golden desensitization."),
        },
        "selfCollisions": collisions,
        "fingerprintImpact": {
            "generatable": generatable,
            "tokenDependentTypes": type_forks,
            "unmatchableInputs": len(rejects),
            "note": ("Rows whose legacy-inferred type is token-dependent "
                     "(installment/fee/interest/cash_advance) risk a fingerprint fork "
                     "against native parser types. Exact proof in web/tests/migration.test.ts."),
        },
    }
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.out:
        Path(args.out).write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0


PAGE_POLLUTION = re.compile(r"第\s*\d+\s*页|共\s*\d+\s*页")
TOKEN_DEPENDENT_TYPES = frozenset({
    "installment_payment", "installment_original", "fee", "interest", "cash_advance",
})


def tag_edge(row: dict) -> str | None:
    tags = []
    if row.get("legacyType") in TOKEN_DEPENDENT_TYPES:
        tags.append("token_dependent_type")
    if PAGE_POLLUTION.search(row.get("description", "")):
        tags.append("boc_page_pollution")
    return "|".join(tags) if tags else None


def load_web_loose_map(golden_root: Path) -> dict[str, list[str]]:
    loose_map: dict[str, list[str]] = {}
    for web_path in sorted(glob.glob(str(golden_root / "*" / "*" / "web.json"))):
        doc = json.loads(Path(web_path).read_text(encoding="utf-8"))
        for tx in doc.get("transactions", []):
            key = "|".join(str(tx.get(k)) for k in (
                "bank", "cardLast4", "transactionDate", "amountMinor", "currency",
            ))
            loose_map.setdefault(key, []).append(tx.get("nativeId", ""))
    return loose_map


def cmd_bridge(args: argparse.Namespace) -> int:
    from ccparser.golden import desensitize_description
    from ccparser.migration_bridge import (
        BRIDGE_SCHEMA,
        build_bridge,
        derive_python_id,
        duplicate_simulation,
    )

    standard_path = Path(args.standard)
    _, legacy_rows, rejects = load_standard_sheet(standard_path)
    for row in legacy_rows:
        row["pythonId"] = derive_python_id(
            row["bankNative"], row["cardLast4"], row["transactionDate"],
            row["postingDate"], row["settlementAmountMinor"], row["description"],
        )
        row["edge"] = tag_edge(row)
    fresh_rows, fresh_missing = load_fresh_parse_rows(Path(args.golden_dir), Path("processed"))
    web_loose_map = load_web_loose_map(Path(args.golden_dir))
    candidates = build_bridge(legacy_rows, fresh_rows, web_loose_map)

    matched = [c for c in candidates if c.matchType == "MATCHED"]
    suspected = [c for c in candidates if c.matchType == "SUSPECTED"]
    orphan = [c for c in candidates if c.matchType == "ORPHAN"]
    simulation = duplicate_simulation(candidates)

    edge_counts: Counter = Counter()
    edge_examples: dict[str, list] = {}
    for row in legacy_rows:
        if row.get("edge"):
            edge_counts[row["edge"]] += 1
            if len(edge_examples.setdefault(row["edge"], [])) < 3:
                edge_examples[row["edge"]].append({
                    "sourceRow": row["sourceRow"],
                    "description": desensitize_description(row["description"])[:80]
                    if not args.raw_details else row["description"][:80],
                })

    def detail(candidate) -> dict:
        # Descriptions stay desensitized in committed output; raw text never leaves.
        return {"pythonId": candidate.pythonId, "candidateWebId": candidate.webId,
                "candidateWebIds": candidate.candidateWebIds, "status": candidate.matchType,
                "confidence": candidate.confidence, "reason": candidate.reason,
                "legacyRow": candidate.legacyRow, "edge": candidate.edge}

    report = {
        "schema": BRIDGE_SCHEMA,
        "scope": {
            "legacyFile": str(standard_path),
            "goldenStatements": sorted(glob.glob(str(Path(args.golden_dir) / "*" / "*"))),
            "freshInputsMissing": fresh_missing,
            "merged": False,
            "note": ("Dry-run only. No database writes, no merges, no deletions. "
                     "ORPHAN means no counterpart inside the 6-statement golden scope, "
                     "not necessarily unmigratable history."),
        },
        "summary": {"total": len(candidates), "matched": len(matched),
                    "suspected": len(suspected), "orphan": len(orphan)},
        "duplicateSimulation": simulation,
        "edgeCases": {
            "counts": dict(edge_counts),
            "examples": edge_examples,
            "rejects": {"total": len(rejects),
                        "reasons": dict(Counter(r["reason"] for r in rejects))},
        },
        "details": [detail(c) for c in candidates],
    }
    text = json.dumps(report, ensure_ascii=False, indent=2)
    out_path = Path(args.out or "reports/migration-report.json")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(text + "\n", encoding="utf-8")
    print(f"total={len(candidates)} matched={len(matched)} "
          f"suspected={len(suspected)} orphan={len(orphan)}")
    print(f"before_duplicates={simulation['before']['duplicatesWouldInsert']} "
          f"after_duplicates={simulation['after']['duplicatesWouldInsert']} "
          f"review_queue={simulation['after']['reviewQueue']}")
    print(f"wrote {out_path} ({out_path.stat().st_size} bytes)")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Legacy migration identity audit (read-only).")
    sub = parser.add_subparsers(dest="command")
    sub.add_parser("audit", help="Inventory + business-key matching report (stdout).")

    bridge = sub.add_parser("bridge", help="Dry-run migration mapping -> migration-report.json.")
    bridge.add_argument("--raw-details", action="store_true",
                        help="Local debugging only: keep raw descriptions (never commit).")
    bridge.add_argument("--standard", default="output/unified_transactions.xlsx")
    bridge.add_argument("--golden-dir", default="golden")
    bridge.add_argument("--out", default=None)

    parser.add_argument("--standard", default="output/unified_transactions.xlsx",
                        help="Legacy 标准字段 export (strict adapter input).")
    parser.add_argument("--golden-dir", default="golden")
    parser.add_argument("--out", default=None, help="Write JSON report to this path.")
    args = parser.parse_args(argv)
    if (args.command or "audit") == "bridge":
        return cmd_bridge(args)
    return cmd_audit(args)


if __name__ == "__main__":
    raise SystemExit(main())
